import os
import shutil
import zipfile
import unittest
import numpy as np
import cv2
import fitz
from core.converter import convert_to_pdf, convert_pages_to_pdf, convert_image_to_pdf

class TestConverter(unittest.TestCase):
    def setUp(self):
        self.test_dir = "spool/test_converter"
        os.makedirs(self.test_dir, exist_ok=True)

    def tearDown(self):
        shutil.rmtree(self.test_dir, ignore_errors=True)

    def test_pdf_passthrough(self):
        doc = fitz.open()
        p = doc.new_page(width=500, height=500)
        p.insert_text((50, 50), "Hello PDF")
        in_path = os.path.join(self.test_dir, "input.pdf")
        doc.save(in_path)
        doc.close()

        out_path = os.path.join(self.test_dir, "output.pdf")
        res = convert_to_pdf(in_path, out_path)
        self.assertTrue(os.path.exists(res))
        doc_out = fitz.open(res)
        self.assertEqual(len(doc_out), 1)
        doc_out.close()

    def test_pages_extraction(self):
        # Create a mock .pages bundle with QuickLook/Preview.pdf
        doc = fitz.open()
        p = doc.new_page(width=500, height=500)
        p.insert_text((50, 50), "Pages Extracted Text")
        pdf_bytes = doc.tobytes()
        doc.close()

        pages_path = os.path.join(self.test_dir, "doc.pages")
        with zipfile.ZipFile(pages_path, "w") as zf:
            zf.writestr("QuickLook/Preview.pdf", pdf_bytes)

        out_path = os.path.join(self.test_dir, "from_pages.pdf")
        res = convert_to_pdf(pages_path, out_path)
        self.assertTrue(os.path.exists(res))
        doc_out = fitz.open(res)
        self.assertIn("Pages Extracted Text", doc_out[0].get_text())
        doc_out.close()

    def test_image_to_pdf(self):
        img = np.full((300, 400, 3), 180, dtype=np.uint8)
        cv2.putText(img, "Test Image", (20, 100), cv2.FONT_HERSHEY_SIMPLEX, 1, (0,0,0), 2)
        img_path = os.path.join(self.test_dir, "test.png")
        cv2.imwrite(img_path, img)

        out_path = os.path.join(self.test_dir, "from_img.pdf")
        res = convert_to_pdf(img_path, out_path)
        self.assertTrue(os.path.exists(res))
        doc_out = fitz.open(res)
        self.assertEqual(len(doc_out), 1)
        doc_out.close()

    def test_docx_conversion(self):
        if os.path.exists("打印机省墨软件开发方案.docx"):
            out_path = os.path.join(self.test_dir, "from_docx.pdf")
            res = convert_to_pdf("打印机省墨软件开发方案.docx", out_path)
            self.assertTrue(os.path.exists(res))
            doc_out = fitz.open(res)
            self.assertGreater(len(doc_out), 0)
            doc_out.close()

    def test_pages_directory_package(self):
        # Create a mock .pages directory package with QuickLook/Preview.pdf
        doc = fitz.open()
        p = doc.new_page(width=500, height=500)
        p.insert_text((50, 50), "Pages Bundle Directory Text")
        pdf_bytes = doc.tobytes()
        doc.close()

        pkg_dir = os.path.join(self.test_dir, "bundle.pages")
        ql_dir = os.path.join(pkg_dir, "QuickLook")
        os.makedirs(ql_dir, exist_ok=True)
        with open(os.path.join(ql_dir, "Preview.pdf"), "wb") as f:
            f.write(pdf_bytes)

        out_path = os.path.join(self.test_dir, "from_pkg.pdf")
        res = convert_to_pdf(pkg_dir, out_path)
        self.assertTrue(os.path.exists(res))
        doc_out = fitz.open(res)
        self.assertIn("Pages Bundle Directory Text", doc_out[0].get_text())
        doc_out.close()

    def test_text_to_pdf(self):
        txt_path = os.path.join(self.test_dir, "sample.txt")
        with open(txt_path, "w", encoding="utf-8") as f:
            f.write("Line 1: Plain text document.\nLine 2: Eco printing rules.\n")

        out_path = os.path.join(self.test_dir, "from_txt.pdf")
        res = convert_to_pdf(txt_path, out_path)
        self.assertTrue(os.path.exists(res))
        doc_out = fitz.open(res)
        self.assertIn("Plain text document", doc_out[0].get_text())
        doc_out.close()

    def test_multipage_pages_ql_bundle(self):
        # Create a mock .pages bundle with 3 distinct pages in Preview.html
        pkg_dir = os.path.join(self.test_dir, "multipage.pages")
        ql_dir = os.path.join(pkg_dir, "QuickLook")
        os.makedirs(ql_dir, exist_ok=True)

        for i, text in enumerate(["Page One Content", "Page Two Content", "Page Three Content"], start=1):
            d = fitz.open()
            p = d.new_page(width=595, height=842)
            p.insert_text((50, 50), text)
            d.save(os.path.join(ql_dir, f"Attachment{i}.pdf"))
            d.close()

        # HTML with reordered sequence: 2, 1, 3
        html_content = (
            '<html><body>'
            '<img src="Attachment2.pdf">'
            '<img src="Attachment1.pdf">'
            '<img src="Attachment3.pdf">'
            '</body></html>'
        )
        with open(os.path.join(ql_dir, "Preview.html"), "w", encoding="utf-8") as f:
            f.write(html_content)

        out_path = os.path.join(self.test_dir, "from_multipage_pages.pdf")
        res = convert_to_pdf(pkg_dir, out_path)
        self.assertTrue(os.path.exists(res))
        doc_out = fitz.open(res)
        self.assertEqual(len(doc_out), 3)
        # Verify page order matches Preview.html (2 -> 1 -> 3)
        self.assertIn("Page Two Content", doc_out[0].get_text())
        self.assertIn("Page One Content", doc_out[1].get_text())
        self.assertIn("Page Three Content", doc_out[2].get_text())
        doc_out.close()

    def test_real_sample_pages_if_exists(self):
        sample = "spool/c832b1b2-40d_orig.pages"
        if os.path.exists(sample) and os.path.exists("/usr/bin/qlmanage"):
            out_path = os.path.join(self.test_dir, "real_pages_out.pdf")
            res = convert_to_pdf(sample, out_path)
            self.assertTrue(os.path.exists(res))
            doc_out = fitz.open(res)
            # Must convert ALL 6 pages, not just 1
            self.assertEqual(len(doc_out), 6)
            doc_out.close()

    def test_image_a4_landscape_fit(self):
        # 2400x1200 wide landscape image
        img = np.full((1200, 2400, 3), 200, dtype=np.uint8)
        cv2.putText(img, "Wide Landscape Test", (100, 300), cv2.FONT_HERSHEY_SIMPLEX, 3, (0,0,0), 5)
        img_path = os.path.join(self.test_dir, "wide.png")
        cv2.imwrite(img_path, img)

        out_path = os.path.join(self.test_dir, "wide.pdf")
        res = convert_to_pdf(img_path, out_path)
        self.assertTrue(os.path.exists(res))
        doc_out = fitz.open(res)
        self.assertEqual(len(doc_out), 1)
        # Should be A4 Landscape: width ~ 841.89 pt, height ~ 595.28 pt
        rect = doc_out[0].rect
        self.assertAlmostEqual(rect.width, 841.89, delta=1.0)
        self.assertAlmostEqual(rect.height, 595.28, delta=1.0)
        doc_out.close()

    def test_image_a4_portrait_fit(self):
        # 1200x2400 tall portrait image
        img = np.full((2400, 1200, 3), 200, dtype=np.uint8)
        cv2.putText(img, "Tall Portrait Test", (100, 300), cv2.FONT_HERSHEY_SIMPLEX, 3, (0,0,0), 5)
        img_path = os.path.join(self.test_dir, "tall.png")
        cv2.imwrite(img_path, img)

        out_path = os.path.join(self.test_dir, "tall.pdf")
        res = convert_to_pdf(img_path, out_path)
        self.assertTrue(os.path.exists(res))
        doc_out = fitz.open(res)
        self.assertEqual(len(doc_out), 1)
        # Should be A4 Portrait: width ~ 595.28 pt, height ~ 841.89 pt
        rect = doc_out[0].rect
        self.assertAlmostEqual(rect.width, 595.28, delta=1.0)
        self.assertAlmostEqual(rect.height, 841.89, delta=1.0)
        doc_out.close()

    def test_transparent_png_to_pdf(self):
        from PIL import Image
        rgba = Image.new("RGBA", (800, 600), (0, 100, 200, 128))
        png_path = os.path.join(self.test_dir, "alpha.png")
        rgba.save(png_path)

        out_path = os.path.join(self.test_dir, "alpha.pdf")
        res = convert_to_pdf(png_path, out_path)
        self.assertTrue(os.path.exists(res))
        doc_out = fitz.open(res)
        self.assertEqual(len(doc_out), 1)
        doc_out.close()

    def test_multipage_pages_zipped_bundle(self):
        # Create a mock .pages file that is a zip archive containing QuickLook attachments
        zip_path = os.path.join(self.test_dir, "zipped_multipage.pages")
        with zipfile.ZipFile(zip_path, "w") as zf:
            for i, text in enumerate(["Zip Page 1", "Zip Page 2", "Zip Page 3"], start=1):
                d = fitz.open()
                p = d.new_page(width=595, height=842)
                p.insert_text((50, 50), text)
                pdf_bytes = d.tobytes()
                d.close()
                zf.writestr(f"QuickLook/Attachment{i}.pdf", pdf_bytes)

            html = '<html><body><img src="Attachment1.pdf"><img src="Attachment2.pdf"><img src="Attachment3.pdf"></body></html>'
            zf.writestr("QuickLook/Preview.html", html)

        out_path = os.path.join(self.test_dir, "from_zipped_pages.pdf")
        res = convert_to_pdf(zip_path, out_path)
        self.assertTrue(os.path.exists(res))
        doc_out = fitz.open(res)
        self.assertEqual(len(doc_out), 3)
        self.assertIn("Zip Page 1", doc_out[0].get_text())
        self.assertIn("Zip Page 2", doc_out[1].get_text())
        self.assertIn("Zip Page 3", doc_out[2].get_text())
        doc_out.close()

    def test_multiframe_tiff_image_to_pdf(self):
        # Create a 2-frame TIFF image
        from PIL import Image
        frame1 = Image.new("RGB", (600, 800), (240, 240, 240))
        frame2 = Image.new("RGB", (800, 600), (220, 220, 220))
        tiff_path = os.path.join(self.test_dir, "multiframe.tiff")
        frame1.save(tiff_path, save_all=True, append_images=[frame2])

        out_path = os.path.join(self.test_dir, "multiframe.pdf")
        res = convert_to_pdf(tiff_path, out_path)
        self.assertTrue(os.path.exists(res))
        doc_out = fitz.open(res)
        self.assertEqual(len(doc_out), 2)
        # First page portrait A4, second page landscape A4
        self.assertAlmostEqual(doc_out[0].rect.width, 595.28, delta=1.0)
        self.assertAlmostEqual(doc_out[0].rect.height, 841.89, delta=1.0)
        self.assertAlmostEqual(doc_out[1].rect.width, 841.89, delta=1.0)
        self.assertAlmostEqual(doc_out[1].rect.height, 595.28, delta=1.0)
        doc_out.close()

    def test_pptx_fallback_dimensions_are_a4_landscape(self):
        from pptx import Presentation
        from pptx.util import Inches
        prs = Presentation()
        prs.slides.add_slide(prs.slide_layouts[0])
        pptx_path = os.path.join(self.test_dir, "test_a4.pptx")
        prs.save(pptx_path)

        from core.converter import convert_pptx_fallback
        out_path = os.path.join(self.test_dir, "pptx_a4_out.pdf")
        res = convert_pptx_fallback(pptx_path, out_path)
        self.assertTrue(res)
        doc_out = fitz.open(out_path)
        self.assertEqual(len(doc_out), 1)
        rect = doc_out[0].rect
        self.assertAlmostEqual(rect.width, 841.89, delta=1.0)
        self.assertAlmostEqual(rect.height, 595.28, delta=1.0)
    def test_resume_docx_fidelity_and_pages_count(self):
        resume_docx = "spool/fa761a3c-104_orig.docx"
        if os.path.exists(resume_docx):
            out_path = os.path.join(self.test_dir, "resume_converted.pdf")
            res = convert_to_pdf(resume_docx, out_path)
            self.assertTrue(os.path.exists(res))
            doc_out = fitz.open(res)
            # The original resume should convert cleanly to <= 4 pages, not 7 pages of bloated duplicate cells
            self.assertLessEqual(len(doc_out), 4)
            # Verify contact info is not repeated multiple times on page 0
            p0_text = doc_out[0].get_text()
            self.assertEqual(p0_text.count("HaoranMa0818@iCloud.com"), 1)
            doc_out.close()

    def test_convert_docx_fallback_merged_cells_no_duplication(self):
        import docx
        from core.converter import convert_docx_fallback
        doc = docx.Document()
        tbl = doc.add_table(rows=2, cols=3)
        # Merge top row across all 3 columns
        tbl.rows[0].cells[0].merge(tbl.rows[0].cells[2])
        tbl.rows[0].cells[0].text = "Single Merged Title"
        tbl.rows[1].cells[0].text = "Col 1"
        tbl.rows[1].cells[1].text = "Col 2"
        tbl.rows[1].cells[2].text = "Col 3"
        doc_path = os.path.join(self.test_dir, "merged.docx")
        doc.save(doc_path)

        out_path = os.path.join(self.test_dir, "merged_out.pdf")
        ok = convert_docx_fallback(doc_path, out_path)
        self.assertTrue(ok)
        pdf = fitz.open(out_path)
        txt = pdf[0].get_text()
        # Must only contain 1 instance of the title, NOT 3 duplicate copies
        self.assertEqual(txt.count("Single Merged Title"), 1)
        pdf.close()


if __name__ == "__main__":
    unittest.main()


